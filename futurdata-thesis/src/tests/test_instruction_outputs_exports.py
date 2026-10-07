"""Parts produced by instructions belong to their parent operation in every format."""
import json
from pathlib import Path

import pytest
from PIL import Image
from docx import Document
from pptx import Presentation

from src.main.models import Diagram, ComponentBox, DiamondStep, ActionCircle, Connection
from src.main.loader_se.disassembly_loader.builder import build_guide
from src.main.loader_se.disassembly_loader.adapter import normalize
from src.main.loader_se.disassembly_loader.validation import rule_action_degree
from src.main.loader_se.disassembly_loader.emitter import guide_to_dict
from src.main.services.document_export_service import DocumentExportService
from src.main.services.txt_exporter.txt_exporter import TXTExporter
from src.main.utils.json_exporter import EnhancedJSONExporter


class UnusedRepository:
    def __getattr__(self, name):
        raise AssertionError(f'Export accessed persistence: {name}')


def instruction_diagram(tmp_path):
    diagram = Diagram()
    nodes = {}
    specs = [
        ('root', ComponentBox, 'Complete Coffee Machine'),
        ('operation', DiamondStep, 'Disassemble the case'),
        ('a', ActionCircle, 'Remove the drip tray by hand'),
        ('b', ActionCircle, 'Unplug the machine then remove the water tank'),
        ('c', ActionCircle, 'Remove the capsule bin by hand'),
        ('tray', ComponentBox, 'Drip Tray & Grid'),
        ('tank', ComponentBox, 'Water Tank'),
        ('bin', ComponentBox, 'Capsule Bin'),
        ('middle', ComponentBox, 'Main Housing'),
        ('next', DiamondStep, 'Remove screws'),
        ('next_action', ActionCircle, 'Unscrew the bottom'),
        ('screws', ComponentBox, 'Security Screws'),
    ]
    for index, (key, kind, text) in enumerate(specs):
        node = kind(index * 100, index * 30)
        node.text = text
        if isinstance(node, ComponentBox):
            image = tmp_path / (key + '.png')
            Image.new('RGB', (30, 20), (index * 15, 40, 70)).save(image)
            # Deliberately mark a continuing component as Leaf: topology wins.
            node.properties.update(name=text, node_type='Root' if key == 'root' else 'Leaf',
                                   image_path=str(image))
        elif isinstance(node, DiamondStep):
            node.name = text
        nodes[key] = node
        diagram.add_shape(node)
    edges = [('root', 'operation'), ('operation', 'a'), ('a', 'b'), ('b', 'c'),
             ('a', 'tray'), ('a', 'tank'), ('b', 'tank'), ('c', 'bin'), ('operation', 'tray'),
             ('c', 'middle'), ('middle', 'next'), ('next', 'next_action'),
             ('next_action', 'screws')]
    for source, target in edges:
        diagram.add_connection(Connection(nodes[source], nodes[target]))
    return diagram


def test_instruction_parts_deduplicate_and_stop_at_next_operation(tmp_path):
    diagram = instruction_diagram(tmp_path)
    snapshot = EnhancedJSONExporter(UnusedRepository()).serialize_active_diagram(diagram)
    source = tmp_path / 'graph.json'
    source.write_text(json.dumps(snapshot), encoding='utf-8')
    guide = build_guide(str(source), include_bom=True)
    first, second = guide.steps
    assert {part.name for part in first.outputs} == {'Drip Tray & Grid', 'Water Tank', 'Capsule Bin'}
    assert len(first.outputs) == 3
    assert [part.name for part in first.continues_as] == ['Main Housing']
    assert second.input == first.continues_as[0]
    assert [part.name for part in second.outputs] == ['Security Screws']
    assert len(first.actions) == 3
    assert next(part for part in first.outputs if part.name == 'Drip Tray & Grid').image_path == str(tmp_path / 'tray.png')
    assert {part.name for part in guide.bill_of_materials} == {
        'Drip Tray & Grid', 'Water Tank', 'Capsule Bin', 'Security Screws'}
    assert not rule_action_degree(normalize(snapshot))
    # A true branch to two successor instructions still warrants a warning.
    snapshot['connections'].append(dict(from_id=snapshot['shapes'][2]['id'],
                                        to_id=snapshot['shapes'][4]['id']))
    assert rule_action_degree(normalize(snapshot))


@pytest.mark.parametrize('format_id', ['html', 'pptx', 'docx', 'md', 'txt'])
def test_every_document_export_includes_instruction_removed_parts(tmp_path, format_id):
    diagram = instruction_diagram(tmp_path)
    before = diagram.to_dict()
    output = Path(DocumentExportService(EnhancedJSONExporter(UnusedRepository())).export(
        diagram, str(tmp_path / format_id / 'guide'), format_id))
    expected = {'Drip Tray & Grid', 'Water Tank', 'Capsule Bin'}
    if format_id == 'html':
        text = output.read_text(encoding='utf-8')
        payload, _ = json.JSONDecoder().raw_decode(text.split('const ir=', 1)[1])
        first = payload['steps'][0]
        assert {part['name'] for part in first['outputs']} == expected
        assert [part['name'] for part in first['continues_as']] == ['Main Housing']
        for part in first['outputs']:
            assert (output.parent / part['image']['path']).is_file()
    elif format_id == 'docx':
        paragraphs = [p.text for p in Document(output).paragraphs]
        removed = next(p for p in paragraphs if p.startswith('Outputs:'))
        assert all(name in removed for name in expected)
        assert 'Security Screws' not in removed and 'Main Housing' not in removed
    elif format_id == 'pptx':
        slides = Presentation(output).slides
        first = next(' '.join(s.text for s in slide.shapes if s.has_text_frame)
                     for slide in slides if any(s.has_text_frame and 'Starting assembly: Complete Coffee Machine' in s.text
                                                for s in slide.shapes))
        assert all(name in first for name in expected)
        assert 'Main Housing' in first and 'Security Screws' not in first
    else:
        text = output.read_text(encoding='utf-8')
        for name in expected:
            assert name.replace('&', chr(92) + '&') in text or name in text
        # Each part appears in the first operation's result section as well as the BoM.
        if format_id == 'txt':
            first = text.split('STEP 1', 1)[1].split('STEP 2', 1)[0]
            removed = first.split('Removed components:', 1)[1].split('Remaining assembly:', 1)[0]
        else:
            first = text.split('## Step 1', 1)[1].split('## Step 2', 1)[0]
            removed = first.split('Parts obtained', 1)[1].split('Remaining assembly:', 1)[0]
        assert all(name in removed or name.replace('&', chr(92) + '&') in removed for name in expected)
        assert 'Security Screws' not in removed and 'Main Housing' not in removed
    assert diagram.to_dict() == before


def test_standalone_txt_copies_instruction_part_and_input_images_without_bom(tmp_path):
    diagram = instruction_diagram(tmp_path)
    snapshot = EnhancedJSONExporter(UnusedRepository()).serialize_active_diagram(diagram)
    source = tmp_path / 'graph.json'; source.write_text(json.dumps(snapshot), encoding='utf-8')
    guide = build_guide(str(source))
    payload = guide_to_dict(guide)
    for index, step in enumerate(payload['steps']):
        step['input']['image'] = {'path': 'root.png' if index == 0 else 'middle.png', 'is_url': False}
        for part in step['outputs']:
            part['image']['path'] = Path(part['image']['path']).name
        for part in step['continues_as']:
            part['image']['path'] = Path(part['image']['path']).name
    ir = tmp_path / 'ir.json'; ir.write_text(json.dumps(payload), encoding='utf-8')
    out = tmp_path / 'published' / 'guide.txt'
    TXTExporter().export(str(ir), str(out))
    for name in ['root.png', 'middle.png', 'tray.png', 'tank.png', 'bin.png', 'screws.png']:
        assert (out.parent / name).read_bytes() == (tmp_path / name).read_bytes()


def branching_diagram(tmp_path):
    diagram = instruction_diagram(tmp_path)
    middle = next(s for s in diagram.shapes if s.text == 'Main Housing')
    # Existing branch is the rear operation; add a second front operation.
    rear = next(s for s in diagram.shapes if s.text == 'Remove screws')
    rear.text = rear.name = 'Removing the rear panel'
    front = DiamondStep(900, 200); front.text = front.name = 'Removing the front panel'
    action_a = ActionCircle(1000, 300); action_a.text = 'Remove the plugs on the MMI'
    action_b = ActionCircle(1000, 500); action_b.text = 'Remove the light guide and the MMI'
    panel = ComponentBox(1200, 500); panel.text = 'Front Panel'
    panel.properties.update(name=panel.text, node_type='Leaf')
    for node in (front, action_a, action_b, panel):
        diagram.add_shape(node)
    for source, target in ((middle, front), (front, action_a), (action_a, action_b), (action_b, panel)):
        diagram.add_connection(Connection(source, target))
    return diagram


def test_every_intermediate_operation_branch_is_visited_once(tmp_path):
    from src.main.loader_se.disassembly_loader.models import DepthSpec, DepthMode
    diagram = branching_diagram(tmp_path)
    exporter = EnhancedJSONExporter(UnusedRepository())
    source = tmp_path / 'graph.json'; source.write_text(json.dumps(exporter.serialize_active_diagram(diagram)))
    guide = build_guide(str(source))
    branches = {step.operation: step for step in guide.steps}
    assert set(branches) == {'Disassemble the case', 'Removing the rear panel', 'Removing the front panel'}
    rear = branches['Removing the rear panel']; front = branches['Removing the front panel']
    assert rear.input == front.input and front.input.name == 'Main Housing'
    assert [a.text for a in front.actions] == ['Remove the plugs on the MMI', 'Remove the light guide and the MMI']
    assert [p.name for p in front.outputs] == ['Front Panel']
    assert [p.name for p in rear.outputs] == ['Security Screws']
    # Coordinate changes must not remove/reorder any branch.
    for shape in diagram.shapes:
        shape.x = -shape.x; shape.y = -shape.y
    source.write_text(json.dumps(exporter.serialize_active_diagram(diagram)))
    assert guide_to_dict(build_guide(str(source)))['steps'] == guide_to_dict(guide)['steps']
    # Explicit depth cuts still suppress every operation below the kept assembly.
    cut = build_guide(str(source), depth=DepthSpec(mode=DepthMode.KEEP_MAIN))
    assert len(cut.steps) == 1
    assert any(p.kept_whole for p in cut.steps[0].outputs)
    # Cyclic malformed graphs must not duplicate operations or loop.
    diagram.add_connection(Connection(next(s for s in diagram.shapes if s.text == 'Removing the front panel'),
                                      next(s for s in diagram.shapes if s.text == 'Main Housing')))
    source.write_text(json.dumps(exporter.serialize_active_diagram(diagram)))
    assert len(build_guide(str(source)).steps) == 3


@pytest.mark.parametrize('format_id', ['html', 'pptx', 'docx', 'md', 'txt'])
def test_every_export_contains_both_operation_branches_and_instructions(tmp_path, format_id):
    diagram = branching_diagram(tmp_path)
    before = diagram.to_dict()
    output = Path(DocumentExportService(EnhancedJSONExporter(UnusedRepository())).export(
        diagram, str(tmp_path / format_id / 'branched'), format_id))
    expected = ['Removing the rear panel', 'Removing the front panel',
                'Remove the plugs on the MMI', 'Remove the light guide and the MMI', 'Front Panel']
    if format_id == 'html':
        text = output.read_text(encoding='utf-8')
        payload, _ = json.JSONDecoder().raw_decode(text.split('const ir=', 1)[1])
        assert len(payload['steps']) == 3
        front = next(s for s in payload['steps'] if s['operation'] == expected[1])
        assert front['input']['name'] == 'Main Housing'
        assert [a['text'] for a in front['actions']] == expected[2:4]
        assert [p['name'] for p in front['outputs']] == ['Front Panel']
    elif format_id == 'pptx':
        text = ' '.join(s.text for slide in Presentation(output).slides for s in slide.shapes if s.has_text_frame)
    elif format_id == 'docx':
        text = ' '.join(p.text for p in Document(output).paragraphs)
    else:
        text = output.read_text(encoding='utf-8')
    assert all(value in text for value in expected)
    assert diagram.to_dict() == before


def test_root_operation_fanout_and_direct_instruction_fanout(tmp_path):
    diagram = branching_diagram(tmp_path)
    root = next(s for s in diagram.shapes if s.text == 'Complete Coffee Machine')
    operation = DiamondStep(1800, 500); operation.text = operation.name = 'Independent root operation'
    first = ActionCircle(1900, 300); first.text = 'First parallel instruction'
    second = ActionCircle(1900, 600); second.text = 'Second parallel instruction'
    part = ComponentBox(2100, 500); part.text = 'Additional Removed Part'
    part.properties.update(name=part.text, node_type='Leaf')
    for node in (operation, first, second, part):
        diagram.add_shape(node)
    for source, target in ((root, operation), (operation, first), (operation, second), (second, part)):
        diagram.add_connection(Connection(source, target))
    snapshot = EnhancedJSONExporter(UnusedRepository()).serialize_active_diagram(diagram)
    source = tmp_path / 'root_branches.json'; source.write_text(json.dumps(snapshot))
    guide = build_guide(str(source))
    assert len(guide.steps) == 4
    operation_step = next(s for s in guide.steps if s.operation == operation.text)
    assert operation_step.input.name == root.text
    assert {a.text for a in operation_step.actions} == {first.text, second.text}
    assert [p.name for p in operation_step.outputs] == [part.text]
