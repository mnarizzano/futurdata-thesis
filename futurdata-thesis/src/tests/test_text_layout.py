"""Regression tests for export typography, overflow, and native editor geometry."""
import re
import tkinter as tk
from pathlib import Path

import pytest
from docx import Document
from pptx import Presentation
from src.main.models import ComponentBox, DiamondStep, ActionCircle, Diagram
from src.main.utils.text_layout import normalize_title, normalize_export_titles, wrap_measured
from src.main.views.canvas_view import DiagramCanvas
from src.main.services.document_export_service import DocumentExportService
from src.main.services.pptx_converter.pptx_core.models import WizardDocument
from src.main.services.pptx_converter.pptx_core.options import ExportOptions
from src.main.services.pptx_converter.pptx_core.exporters.pptx_exporter import PPTXExporter
from src.main.utils.json_exporter import EnhancedJSONExporter
from src.tests.test_export_services import export_diagram, UnusedRepository


@pytest.mark.parametrize('value,expected', [
    ('customer information', 'Customer information'),
    ('  123 - tool API configuration', '  123 - Tool API configuration'),
    ('API Configuration', 'API Configuration'),
    ('\u00e9quipement ABC', '\u00c9quipement ABC'), ('123!?', '123!?'), ('', ''),
])
def test_title_changes_only_first_letter(value, expected):
    assert normalize_title(value) == expected


def test_normalization_does_not_modify_source_or_body():
    source = {'product': {'name': 'lower API', 'image': {'path': 'lower.png'}},
              'steps': [{'operation': 'remove ABC', 'actions': [{'text': 'keep Body'}]}]}
    result = normalize_export_titles(source)
    assert result['product']['name'] == 'Lower API'
    assert source['product']['name'] == 'lower API'
    assert result['product']['image']['path'] == 'lower.png'
    assert result['steps'][0]['actions'][0]['text'] == 'keep Body'


def test_wrapping_breaks_long_tokens_without_losing_characters():
    value = 'normal words ' + 'https://example.test/' + 'W' * 500
    lines = wrap_measured(value, 40, len)
    assert ''.join(lines) == value
    assert all(len(line) <= 40 for line in lines)


@pytest.mark.parametrize('format_id', ['html', 'md', 'txt', 'docx', 'pptx'])
def test_all_exports_normalize_titles_without_changing_diagram(tmp_path, export_diagram, format_id):
    diagram, _ = export_diagram
    product, _, operation, _ = diagram.shapes[:4]
    product.text = product.properties['name'] = 'customer API information'
    operation.text = operation.name = 'tool API configuration'
    output = Path(DocumentExportService(EnhancedJSONExporter(UnusedRepository())).export(
        diagram, str(tmp_path / 'result'), format_id))
    if format_id == 'pptx':
        text = ' '.join(s.text for slide in Presentation(output).slides for s in slide.shapes if s.has_text_frame)
    elif format_id == 'docx':
        text = ' '.join(p.text for p in Document(output).paragraphs)
    else:
        text = output.read_text(encoding='utf-8')
    assert 'Customer API information' in text
    assert 'Tool API configuration' in text
    assert product.text == 'customer API information'
    assert operation.name == 'tool API configuration'


@pytest.mark.parametrize('kind', [ComponentBox, DiamondStep, ActionCircle])
def test_native_canvas_grows_and_keeps_text_inside_shape(kind):
    root = tk.Tk()
    root.withdraw()
    try:
        canvas = DiagramCanvas(root)
        shape = kind(300, 300)
        shape.text = 'Long URL https://example.test/' + 'W' * 250
        neighbor = ComponentBox(300, 430)
        diagram = Diagram()
        diagram.shapes.extend([shape, neighbor])
        canvas.redraw_all(diagram)
        left, top, right, bottom = canvas.bbox(shape.text_id)
        # All label corners must lie within curved as well as rectangular shapes.
        for x, y in ((left, top), (right, top), (left, bottom), (right, bottom)):
            assert shape.contains_point(x, y)
        assert neighbor.get_bounds()[1] >= shape.get_bounds()[3] + 20
        assert shape._display_text.replace('\n', '') == shape.text
        positions = [(s.x, s.y) for s in diagram.shapes]
        canvas.redraw_all(diagram)
        assert positions == [(s.x, s.y) for s in diagram.shapes]
        for zoom in (0.4, 1.0, 2.0):
            canvas._apply_zoom(zoom / canvas.zoom_factor)
            bounds = canvas.bbox(shape.text_id)
            for x, y in ((bounds[0], bounds[1]), (bounds[2], bounds[3])):
                assert shape.contains_point(x / zoom, y / zoom)
    finally:
        root.destroy()


@pytest.mark.parametrize('groups', [1, 4])
def test_pptx_overflow_keeps_all_text_in_readable_continuations(tmp_path, groups):
    long_action = 'BODY_START ' + 'W' * 2500 + ' BODY_END'
    tools = ['tool_' + str(i) for i in range(35)]
    outputs = [{'node_id': i + 2, 'name': 'part_' + str(i)} for i in range(12)]
    data = {'product': {'name': 'product API'}, 'steps': [
        {'index': 1, 'operation': 'operation API', 'tools_required': tools,
         'actions': [{'text': long_action}], 'outputs': outputs}],
        'bill_of_materials': [{'node_id': 99, 'name': 'bom_' + 'Z' * 1500 + '_end'}]}
    exporter = PPTXExporter()
    path = exporter.export(WizardDocument.from_any(data), tmp_path / 'long.pptx', ExportOptions(groups_per_slide=groups))
    presentation = Presentation(path)
    texts = []
    for slide in presentation.slides:
        for shape in slide.shapes:
            assert shape.left >= 0 and shape.top >= 0
            assert shape.left + shape.width <= presentation.slide_width + 100
            assert shape.top + shape.height <= presentation.slide_height + 100
            if shape.has_text_frame:
                texts.append(shape.text)
                for paragraph in shape.text_frame.paragraphs:
                    for run in paragraph.runs:
                        assert run.font.size.pt >= 8
            if shape.has_table:
                texts.extend(cell.text for row in shape.table.rows for cell in row.cells)
    text = re.sub(r'\s+', '', ''.join(texts))
    # The content may span slides; headers/footers separate chunks, so test the
    # complete retained original via the pagination queue as well as sentinels.
    assert any(value == long_action for _, _, value in exporter._overflow)
    assert 'BODY_START' in text and 'BODY_END' in text
    assert 'tool_34' in text and 'Part_11' in text and '_end' in text
    assert 'Details[' in text


def test_docx_table_wraps_without_changing_long_url(tmp_path, export_diagram):
    diagram, _ = export_diagram
    long_name = 'part' + 'W' * 500
    diagram.shapes[3].text = long_name
    diagram.shapes[3].properties['name'] = long_name
    path = DocumentExportService(EnhancedJSONExporter(UnusedRepository())).export(diagram, str(tmp_path / 'word'), 'docx')
    doc = Document(path)
    assert doc.tables and all(not table.autofit for table in doc.tables)
    cells = [cell for table in doc.tables for row in table.rows for cell in row.cells]
    assert any(cell.text == normalize_title(long_name) for cell in cells)
    assert all('wordWrap' in p._p.xml for cell in cells for p in cell.paragraphs)


def test_custom_tool_persists_and_clears_through_existing_repository(tmp_path):
    from src.main.controllers.app_controller import AppController
    from src.main.repositories.json_repository import JsonRepository
    controller = AppController.__new__(AppController)
    controller.repository = JsonRepository(str(tmp_path / 'store.json'))
    from src.main.models import Diagram, ComponentBox
    controller.diagram = Diagram()
    root = ComponentBox(0, 0)
    root.properties['node_type'] = 'Root'
    controller.diagram.shapes.append(root)
    controller.current_product_id = None
    shape = DiamondStep(0, 0)
    shape.tools = 'Custom API function'
    controller._persist_shape_properties(shape)
    saved_id = shape.tool_id
    assert saved_id is not None
    assert controller.repository.get_action(shape.db_action_id)['tool_name'] == shape.tools
    controller._persist_shape_properties(shape)
    assert shape.tool_id == saved_id
    shape.tools = ''
    controller._persist_shape_properties(shape)
    assert shape.tool_id == saved_id  # Legacy ID-only diagrams remain valid.
    shape.tool_id = None
    controller._persist_shape_properties(shape)
    assert controller.repository.get_action(shape.db_action_id)['tool_id'] is None


def test_markdown_long_labels_keep_breaks_out_of_link_targets():
    from src.main.services.md_exporter.md_utils import image, escape_cell
    label = 'W' * 200
    assert '<wbr>' in escape_cell(label)
    path = 'https://example.test/' + 'a' * 200
    assert image(label, path).endswith('](' + path + ')')
