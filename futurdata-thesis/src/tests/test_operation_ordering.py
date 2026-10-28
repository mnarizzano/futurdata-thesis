"""Operation-level breadth-first ordering and the supported Load interface."""
import copy
import json
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import MagicMock, patch

import pytest
from docx import Document
from pptx import Presentation

from src.main.loader_se.disassembly_loader import build_guide, DepthMode, DepthSpec
from src.main.services.document_export_service import DocumentExportService
from src.main.views.main_window import MainWindow


def graph(branches=True):
    specs = [(1, "component", "Root"), (10, "diamond", "Operation A"),
             (20, "diamond", "Operation B"), (30, "diamond", "Operation C"),
             (11, "action", "Instruction one"), (12, "action", "Instruction two"),
             (13, "component", "Removed part")]
    edges = [(1, 30), (1, 20), (1, 10), (10, 11), (11, 12), (12, 13)]
    if branches:
        specs += [(40, "component", "Intermediate A"), (50, "component", "Intermediate B"),
                  (60, "diamond", "Operation D"), (70, "diamond", "Operation E")]
        edges += [(10, 40), (40, 60), (20, 50), (50, 70)]
    return {"shapes": [dict(id=i, type=t, text=text, x=-i, y=i,
                             image_path={1: "root.jpg", 40: "middle.jpg", 11: "action.jpg"}.get(i))
                       for i, t, text in reversed(specs)],
            "connections": [dict(from_id=a, to_id=b) for a, b in edges]}


def load(tmp_path, data, **options):
    source = tmp_path / "graph.json"
    source.write_text(json.dumps(data), encoding="utf-8")
    return build_guide(str(source), **options)


def names(guide):
    return [step.operation for step in guide.steps]


def test_same_level_and_branch_levels_preserve_instructions_images_and_outputs(tmp_path):
    assert names(load(tmp_path, graph(False))) == ["Operation A", "Operation B", "Operation C"]
    data = graph()
    before = copy.deepcopy(data)
    guide = load(tmp_path, data)
    assert names(guide) == ["Operation A", "Operation B", "Operation C", "Operation D", "Operation E"]
    assert [s.index for s in guide.steps] == [1, 2, 3, 4, 5]
    assert [s.input.node_id for s in guide.steps] == [1, 1, 1, 40, 50]
    assert [s.input.image_path for s in guide.steps[:3]] == ["root.jpg"] * 3
    assert guide.steps[3].input.image_path == "middle.jpg"
    assert [a.text for a in guide.steps[0].actions] == ["Instruction one", "Instruction two"]
    assert guide.steps[0].actions[0].image_path == "action.jpg"
    assert [p.node_id for p in guide.steps[0].outputs] == [13]
    assert [p.node_id for p in guide.steps[0].continues_as] == [40]
    assert data == before
    data["shapes"].reverse()
    data["connections"].reverse()
    assert load(tmp_path, data).steps == guide.steps
    assert names(load(tmp_path, data, ordering="dfs")) == ["Operation A", "Operation D", "Operation B", "Operation E", "Operation C"]
    cut = load(tmp_path, data, depth=DepthSpec(DepthMode.KEEP_MAIN))
    assert len(cut.steps) == 3
    assert cut.steps[0].outputs[-1].kept_whole


def test_converging_paths_wait_for_all_producers_and_multiple_inputs(tmp_path):
    data = graph()
    data["connections"] += [dict(from_id=70, to_id=40)]
    # D must wait for E, even though its assembly is first reached through A.
    assert names(load(tmp_path, data)) == ["Operation A", "Operation B", "Operation C", "Operation E", "Operation D"]
    data = graph()
    data["shapes"] += [dict(id=80, type="component", text="Later assembly")]
    data["connections"] += [dict(from_id=70, to_id=80), dict(from_id=80, to_id=60)]
    assert names(load(tmp_path, data))[-2:] == ["Operation E", "Operation D"]


def test_disconnected_dangling_and_cyclic_graphs_are_bounded(tmp_path):
    data = graph()
    data["shapes"] += [dict(id=900, type="diamond", text="Disconnected")]
    data["connections"] += [dict(from_id=60, to_id=40), dict(from_id=10, to_id=999)]
    guide = load(tmp_path, data)
    assert len(guide.steps) == 5
    assert len(set(names(guide))) == 5
    assert guide.warnings
    assert names(load(tmp_path, data)) == names(guide)
    data["connections"] += [dict(from_id=60, to_id=1)]
    assert load(tmp_path, data).steps == ()  # existing rootless-graph contract


@pytest.mark.parametrize("fmt", ["html", "pptx", "docx", "md", "txt"])
def test_every_export_consumes_canonical_operation_order(tmp_path, fmt):
    data = graph()
    before = copy.deepcopy(data)
    exporter = SimpleNamespace(serialize_active_diagram=lambda diagram: copy.deepcopy(data))
    path = DocumentExportService(exporter).export(SimpleNamespace(file_path=None), str(tmp_path / fmt / "guide"), fmt)
    if fmt == "html":
        text = Path(path).read_text(encoding="utf-8")
        payload, _ = json.JSONDecoder().raw_decode(text.split("const ir=", 1)[1])
        assert [step["operation"] for step in payload["steps"]] == names(load(tmp_path, data))
        assert [step["input"]["node_id"] for step in payload["steps"]] == [1, 1, 1, 40, 50]
    else:
        if fmt == "pptx":
            text = "\n".join(shape.text for slide in Presentation(path).slides for shape in slide.shapes if shape.has_text_frame)
        elif fmt == "docx":
            text = "\n".join(p.text for p in Document(path).paragraphs)
        else:
            text = Path(path).read_text(encoding="utf-8")
        positions = [text.index(title) for title in names(load(tmp_path, data))]
        assert positions == sorted(positions)
    assert data == before
