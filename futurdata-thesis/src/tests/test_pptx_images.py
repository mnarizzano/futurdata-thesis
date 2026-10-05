"""Exercise picture insertion through real PowerPoint image relationships."""

import base64
from io import BytesIO

from PIL import Image
from pptx import Presentation
import pytest

from src.main.services.pptx_converter.pptx_core.exporters.pptx_exporter import (
    PPTXExporter,
)


@pytest.mark.parametrize("kind", ["relative", "absolute", "data_uri", "remote"])
def test_picture_supports_local_paths_and_streams(tmp_path, monkeypatch, kind):
    image = tmp_path / "part with spaces.png"
    Image.new("RGB", (80, 40), "red").save(image)
    payload = image.read_bytes()
    value = image.name if kind == "relative" else str(image)
    if kind == "data_uri":
        value = "data:image/png;base64," + base64.b64encode(payload).decode("ascii")
    elif kind == "remote":
        value = "https://example.invalid/part.png"
        monkeypatch.setattr(
            "src.main.services.pptx_converter.pptx_core.utils._download_image",
            lambda url: BytesIO(payload),
        )
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    assert PPTXExporter()._picture(slide, value, str(tmp_path), 1, 1, 4, 4)
    saved = tmp_path / "picture.pptx"
    presentation.save(str(saved))
    picture = Presentation(str(saved)).slides[0].shapes[0]
    assert picture.image.blob == payload
    assert picture.width / picture.height == pytest.approx(2)


@pytest.mark.parametrize("kind", ["missing", "corrupt"])
def test_unavailable_picture_remains_optional(tmp_path, kind):
    path = tmp_path / "unavailable.png"
    if kind == "corrupt":
        path.write_bytes(b"not an image")
    presentation = Presentation()
    slide = presentation.slides.add_slide(presentation.slide_layouts[6])
    assert not PPTXExporter()._picture(slide, str(path), str(tmp_path), 1, 1, 4, 4)
    assert len(slide.shapes) == 0
