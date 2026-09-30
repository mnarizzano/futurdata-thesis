"""Exercise desktop JSON import through the actual presentation service."""
import json
from pathlib import Path

import pytest
from PIL import Image
from pptx import Presentation

from src.main.models import ComponentBox
from src.main.services.document_export_service import DocumentExportService
from src.main.utils.image_handler import ImageHandler
from src.main.utils.json_exporter import EnhancedJSONExporter


@pytest.mark.parametrize("keys", [("from_id", "to_id"), ("from_shape_id", "to_shape_id")])
@pytest.mark.parametrize("image_folder", ["images", "diagram_images"])
def test_import_connections_root_and_images_export_to_pptx(tmp_path, monkeypatch, keys, image_folder):
    handler = ImageHandler(str(tmp_path / "store"))
    monkeypatch.setattr("src.main.utils.image_handler.get_image_handler", lambda: handler)
    monkeypatch.setattr("src.main.services.presentation_service.get_image_handler", lambda: handler)
    folder = tmp_path / image_folder
    folder.mkdir()
    image = folder / "page.png"
    Image.new("RGB", (40, 60), "purple").save(image)
    source, target = keys
    shapes = [
        dict(id=1, type="product", x=100, y=100, text="Washer", name="Washer"),
        dict(id=2, type="diamond", x=100, y=300, text="Remove cover", name="Remove cover"),
        dict(id=3, type="action", x=400, y=300, text="1.1", step_description="Release cover", image_path="images/page.png"),
        dict(id=4, type="component", x=100, y=500, text="Cover", node_type="Leaf"),
    ]
    edges = [{source: a, target: b, "from_anchor": "right", "to_anchor": "left"} for a,b in [(1,2),(2,3),(2,4)]]
    data = dict(metadata={"description": "source notes"}, diagram={"canvas_size": [3400,24000]}, shapes=shapes, connections=edges)
    path = tmp_path / "diagram.json"
    path.write_text(json.dumps(data), encoding="utf-8")
    exporter = EnhancedJSONExporter(None)
    diagram = exporter.import_diagram(str(path), create_in_repository=False)
    assert diagram is not None
    assert len(diagram.connections) == 3
    assert diagram.connections[0].from_anchor == "right"
    assert diagram.canvas_size == (3400,24000)
    assert diagram.metadata["description"] == "source notes"
    assert isinstance(diagram.shapes[0], ComponentBox)
    assert diagram.shapes[0].properties["node_type"] == "Root"
    assert Path(handler.get_full_path(diagram.shapes[2].image_path)).is_file()
    output = DocumentExportService(exporter).export(diagram, str(tmp_path / "guide.pptx"), "pptx")
    presentation = Presentation(output)
    text = "\n".join(s.text for slide in presentation.slides for s in slide.shapes if s.has_text_frame)
    assert "Remove cover" in text
    assert "Release cover" in text
    pictures = [s.image.blob for slide in presentation.slides for s in slide.shapes if s.shape_type == 13]
    assert image.read_bytes() in pictures
