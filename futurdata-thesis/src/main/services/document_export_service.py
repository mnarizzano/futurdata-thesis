"""Unified exports for the active diagram: PPTX, DOCX, Markdown, TXT and HTML."""

from __future__ import annotations
import json
import shutil
import tempfile
from pathlib import Path
from ..utils.material_display import material_text
from ..utils.image_handler import get_image_handler
from ..utils.text_layout import normalize_export_titles
from ..loader_se.disassembly_loader import build_guide, write_json


class DocumentExportService:
    FORMATS = {
        "pptx": ("PowerPoint presentation", ".pptx"),
        "docx": ("Word document", ".docx"),
        "md": ("Markdown", ".md"),
        "txt": ("Text", ".txt"),
        "html": ("HTML/Web", ".html"),
    }

    def __init__(self, diagram_exporter):
        self.diagram_exporter = diagram_exporter

    def export(self, diagram, output_path: str, format_id: str) -> str:
        fmt = format_id.lower()
        if fmt not in self.FORMATS:
            raise ValueError(f"Unsupported export format: {fmt}")
        dest = Path(output_path).expanduser().with_suffix(self.FORMATS[fmt][1])
        dest.parent.mkdir(parents=True, exist_ok=True)
        if fmt == "pptx":
            from .presentation_service import PresentationExportService

            return PresentationExportService(self.diagram_exporter).export(
                diagram, str(dest)
            )
        snapshot = self.diagram_exporter.serialize_active_diagram(diagram)
        with tempfile.TemporaryDirectory(prefix=f"ariadne_{fmt}_") as td:
            root = Path(td)
            self._stage_images(snapshot, root)
            source = root / "diagram.json"
            source.write_text(
                json.dumps(snapshot, indent=2, ensure_ascii=False), encoding="utf-8"
            )
            guide = build_guide(str(source), include_bom=True)
            ir = root / "ir.json"
            write_json(guide, str(ir))
            # Text/web formats reference image files rather than embedding them.
            # Copy the staged assets next to the final document before the
            # temporary export workspace is removed.
            if fmt in {"html", "txt", "md"}:
                self._publish_images(root / "images", dest.parent / "images")
            if fmt == "html":
                from .html_exporter.html_exporter import HTMLExporter

                HTMLExporter().export(str(ir), str(dest))
            elif fmt == "txt":
                from .txt_exporter.txt_exporter import TXTExporter

                TXTExporter().export(str(ir), str(dest))
            elif fmt == "md":
                from .md_exporter.exporter import MDExporter, ExportOptions

                MDExporter().export(guide, str(dest), ExportOptions())
            elif fmt == "docx":
                self._export_docx(guide, dest, root)
        return str(dest.resolve())

    @staticmethod
    def _stage_images(snapshot, root):
        from .image_staging import stage_images
        stage_images(snapshot, root, get_image_handler())

    @staticmethod
    def _publish_images(source_dir, destination_dir):
        if not source_dir.is_dir():
            return
        for source in source_dir.rglob("*"):
            if not source.is_file():
                continue
            target = destination_dir / source.relative_to(source_dir)
            target.parent.mkdir(parents=True, exist_ok=True)
            shutil.copy2(source, target)

    @staticmethod
    def _export_docx(guide, dest, source_root):
        from docx import Document

        guide = normalize_export_titles(guide)
        doc = Document()
        doc.add_heading(f"Disassembly guide: {guide.product.name}", 0)

        def add_image(image_path, width_inches=5.5):
            if not image_path or str(image_path).startswith(("http://", "https://")):
                return
            path = Path(image_path)
            if not path.is_absolute():
                path = source_root / path
            if path.is_file():
                try:
                    from docx.shared import Inches

                    doc.add_picture(str(path), width=Inches(width_inches))
                except (OSError, ValueError):
                    pass

        add_image(getattr(guide.product, "image_path", None))
        if guide.warnings:
            doc.add_heading("Validation warnings", 1)
            for w in guide.warnings:
                doc.add_paragraph(f"[{w.severity.value.upper()}] {w.message}")
        if guide.bill_of_materials:
            doc.add_heading("Bill of materials", 1)
            table = doc.add_table(rows=1, cols=4)
            table.style = "Table Grid"
            for c, t in zip(
                table.rows[0].cells, ["Component", "Weight", "Material", "Color"]
            ):
                c.text = t
            for item in guide.bill_of_materials:
                cells = table.add_row().cells
                cells[0].text = str(item.name)
                cells[1].text = (
                    f"{item.weight} {item.weight_unit or ''}"
                    if item.weight is not None
                    else "—"
                )
                cells[2].text = material_text(item.material, item.material_details) or "Unknown"
                cells[3].text = str(item.color or "—")
        for step in guide.steps:
            doc.add_heading(f"Step {step.index}: {step.operation}", 1)
            doc.add_paragraph(f"Input: {step.input.name}")
            if step.tools_required:
                doc.add_paragraph("Tools required: " + ", ".join(step.tools_required))
            for action in step.actions:
                doc.add_paragraph(action.text, style="List Bullet")
                add_image(getattr(action, "image_path", None), 4.5)
            if step.outputs:
                doc.add_paragraph("Outputs: " + ", ".join(x.name for x in step.outputs))
                for output in step.outputs:
                    add_image(getattr(output, "image_path", None), 4.5)
            if step.continues_as:
                doc.add_paragraph(
                    "Continue with: " + ", ".join(x.name for x in step.continues_as)
                )
        from docx.oxml import OxmlElement
        from docx.oxml.ns import qn
        from docx.enum.table import WD_ROW_HEIGHT_RULE
        # Permit breaking Latin words/URLs without inserting data characters.
        paragraphs = list(doc.paragraphs)
        usable_width = doc.sections[0].page_width - doc.sections[0].left_margin - doc.sections[0].right_margin
        for table in doc.tables:
            table.autofit = False
            width = int(usable_width / len(table.columns))
            for column in table.columns:
                column.width = width
            for row in table.rows:
                row.height_rule = WD_ROW_HEIGHT_RULE.AUTO
                for cell in row.cells:
                    cell.width = width
                    paragraphs.extend(cell.paragraphs)
        for paragraph in paragraphs:
            wrap = OxmlElement('w:wordWrap')
            wrap.set(qn('w:val'), '0')
            paragraph._p.get_or_add_pPr().append(wrap)
        doc.save(str(dest))
