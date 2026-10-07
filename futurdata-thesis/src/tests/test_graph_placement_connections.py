"""Graph semantics and actual Tk item ownership, independent of screenshot layout."""
import json
import tkinter as tk
from types import SimpleNamespace
from unittest.mock import MagicMock
import pytest
from PIL import Image
from pptx import Presentation
from pptx.enum.shapes import MSO_SHAPE_TYPE
from src.main.models import Diagram, ComponentBox, DiamondStep, ActionCircle, ArrowShape, Connection
from src.main.views.canvas_view import DiagramCanvas
from src.main.utils.commands import MoveShapeCommand, CommandHistory
from src.main.utils.json_exporter import EnhancedJSONExporter
from src.main.loader_se.disassembly_loader.builder import build_guide
from src.main.loader_se.disassembly_loader.adapter import normalize
from src.main.loader_se.disassembly_loader.emitter import guide_to_dict
from src.main.services.pptx_converter.pptx_core.models import WizardDocument
from src.main.services.pptx_converter.pptx_core.exporters.pptx_exporter import PPTXExporter
from src.main.services.pptx_converter.pptx_core.options import ExportOptions
from src.main.services.html_exporter.scripts import get_scripts


def graph_payload():
    shapes = []
    for i, role, name, image in [(1,'component','Root','root.png'),
                                (2,'diamond','Remove outer parts','operation.png'),
                                (3,'component','Leaf A','a.png'),
                                (4,'component','Leaf B',None),
                                (5,'component','Middle','middle.png'),
                                (6,'diamond','Remove screws',None),
                                (7,'action','Instruction A','act_a.png'),
                                (8,'action','Instruction B','act_b.png'),
                                (9,'component','Screws',None)]:
        shapes.append(dict(id=i,type=role,text=name,image_path=image,x=i*100,y=0))
    edges=[(1,2),(2,3),(2,4),(2,5),(5,6),(2,7),(7,8),(6,9)]
    return dict(shapes=shapes,connections=[dict(from_id=a,to_id=b) for a,b in edges])


def test_canonical_step_roles_and_media_ignore_coordinates(tmp_path):
    data=graph_payload()
    path=tmp_path/'graph.json';path.write_text(json.dumps(data))
    guide=build_guide(str(path))
    first,second=guide.steps
    assert first.input.name=='Root' and first.input.image_path=='root.png'
    assert [p.name for p in first.outputs]==['Leaf A','Leaf B']
    assert [p.name for p in first.continues_as]==['Middle']
    assert second.input.name=='Middle' and second.input.image_path=='middle.png'
    assert [a.text for a in first.actions]==['Instruction A','Instruction B']
    assert [a.image_path for a in first.actions]==['act_a.png','act_b.png']
    assert first.image_path=='operation.png'
    before=guide_to_dict(guide)
    for shape in data['shapes']:
        shape['x']=-shape['x'];shape['y']=99999-shape['id']
        shape['node_type']='Leaf' # metadata does not decide continuation
    path.write_text(json.dumps(data))
    assert guide_to_dict(build_guide(str(path)))==before
    document=WizardDocument.from_any(before)
    assert document.steps[1].source.image=='middle.png'
    assert [p.name for p in document.steps[0].outputs]==['Leaf A','Leaf B']
    assert document.steps[0].actions[1].image=='act_b.png'
    import jsonschema
    from pathlib import Path
    schema=Path(__file__).parents[1]/'main/loader_se/disassembly_loader/ir_schema.json'
    jsonschema.validate(before,json.loads(schema.read_text()))


def test_mixed_import_has_one_edge_and_keeps_arrow_only_edges(tmp_path):
    diagram=Diagram();a=ComponentBox(100,100);b=DiamondStep(300,300);c=ActionCircle(500,500)
    diagram.shapes=[a,b,c,ArrowShape(0,0,a,b),ArrowShape(0,0,b,c)]
    snapshot=EnhancedJSONExporter(MagicMock()).serialize_active_diagram(diagram)
    path=tmp_path/'snapshot.json';path.write_text(json.dumps(snapshot))
    imported=EnhancedJSONExporter(MagicMock()).import_diagram(str(path),False)
    assert imported is not None
    assert len(imported.connections)+sum(isinstance(s,ArrowShape) for s in imported.shapes)==2
    native=diagram.to_dict();native['connections']=[Connection(a,b).to_dict()]*2
    restored=Diagram.from_dict(native)
    assert len(restored.connections)==0
    restored.add_connection(Connection(restored.shapes[0],restored.shapes[1]))
    assert len(restored.connections)==0
    data=graph_payload()
    data['shapes'].append(dict(type='arrow',from_shape_id=6,to_shape_id=7))
    graph=normalize(data)
    assert (6,7) in [(e.src,e.dst) for e in graph.edges]


@pytest.fixture
def canvas():
    root=tk.Tk();root.geometry('900x700')
    view=DiagramCanvas(root);view.pack(fill='both',expand=True);root.update()
    yield view,root
    root.destroy()


@pytest.mark.parametrize('legacy',[False,True])
def test_one_rendered_edge_through_moves_zoom_redraw_and_load(canvas,legacy):
    view,root=canvas
    diagram=Diagram();a=ComponentBox(250,250);b=DiamondStep(550,500);diagram.shapes=[a,b]
    edge=ArrowShape(0,0,a,b) if legacy else Connection(a,b)
    if legacy: diagram.add_shape(edge)
    else: diagram.add_connection(edge)
    history=CommandHistory();view.redraw_all(diagram)
    for i in range(8):
        item=view._canvas_items[edge]['body'] if legacy else view._connection_items[edge]
        node=a if i%2 else b
        history.execute(MoveShapeCommand(node,15,10,diagram))
        view.move_items(node,15,10);view.update_connections_for_shapes([node],diagram)
        assert item==(view._canvas_items[edge]['body'] if legacy else view._connection_items[edge])
        diagram.select_shape(node);view.redraw_all(diagram);diagram.clear_selection();view.redraw_all(diagram)
        view._apply_zoom(1.05,100,100);view.redraw_all(diagram)
        arrows=[i for i in view.find_all() if view.type(i)=='line' and view.itemcget(i,'arrow')!='none']
        assert len(arrows)==1
        assert len(diagram.connections)+sum(isinstance(s,ArrowShape) for s in diagram.shapes)==1
    history.undo();view.redraw_all(diagram);history.redo();view.redraw_all(diagram)
    restored=Diagram.from_dict(diagram.to_dict());view.redraw_all(restored)
    assert sum(view.type(i)=='line' and view.itemcget(i,'arrow')!='none' for i in view.find_all())==1


@pytest.mark.parametrize('zoom',[.5,1,2])
def test_placement_scrolled_zoomed_occupied_viewport(canvas,zoom):
    view,root=canvas
    view._apply_zoom(zoom,0,0);view.xview_moveto(.2);view.yview_moveto(.15);root.update()
    cx,cy=view.model_x(view.winfo_width()/2),view.model_y(view.winfo_height()/2)
    diagram=Diagram();center=ComponentBox(cx,cy);diagram.add_shape(center);view.redraw_all(diagram);root.update()
    # Redraw can expand scrollregion, so obtain the actual visible center again.
    center.x=view.model_x(view.winfo_width()/2);center.y=view.model_y(view.winfo_height()/2)
    view.redraw_all(diagram);before=diagram.to_dict()
    shape=ComponentBox(0,0);x,y=view.find_free_position(shape,diagram)
    shape.x,shape.y=x,y
    l,t,r,b=shape.get_bounds();ol,ot,ore,ob=view.render_bounds(center)
    assert r<=ol or l>=ore or b<=ot or t>=ob
    assert l>=view.model_x(0) and r<=view.model_x(view.winfo_width())
    assert t>=view.model_y(0) and b<=view.model_y(view.winfo_height())
    assert diagram.to_dict()==before


def test_pptx_main_image_matches_input_and_all_instruction_media_survive(tmp_path):
    data=graph_payload();path=tmp_path/'graph.json';path.write_text(json.dumps(data))
    guide=build_guide(str(path));payload=guide_to_dict(guide)
    expected={}
    for i,name in enumerate(['root.png','middle.png','operation.png','a.png','act_a.png','act_b.png']):
        image=tmp_path/name;Image.new('RGB',(30,20),(i*35,20,50)).save(image);expected[name]=image.read_bytes()
    document=WizardDocument.from_any(payload,source_dir=str(tmp_path))
    out=tmp_path/'guide.pptx';PPTXExporter().export(document,out,ExportOptions())
    prs=Presentation(str(out));pictures=[]
    for slide in prs.slides:
        text=' '.join(s.text for s in slide.shapes if s.has_text_frame)
        blobs=[s.image.blob for s in slide.shapes if s.shape_type==MSO_SHAPE_TYPE.PICTURE]
        pictures.extend(blobs)
        if 'Starting assembly: Middle' in text:
            assert expected['middle.png'] in blobs
            assert expected['root.png'] not in blobs
    assert set(expected.values())<=set(pictures)
    scripts=get_scripts()
    assert 'const p=s.input?.image?.path' in scripts
    assert 'firstImage' not in scripts
    assert 'Array.isArray(c)' in scripts


def test_html_copies_input_and_operation_assets_without_bom(tmp_path):
    from src.main.services.html_exporter.html_exporter import HTMLExporter
    data=graph_payload();source=tmp_path/'source';source.mkdir()
    path=source/'graph.json';path.write_text(json.dumps(data))
    ir=guide_to_dict(build_guide(str(path)))
    for name in ['root.png','middle.png','operation.png','a.png','act_a.png','act_b.png']:
        Image.new('RGB',(20,10),'red').save(source/name)
    ir_path=source/'ir.json';ir_path.write_text(json.dumps(ir))
    out=tmp_path/'published'/'guide.html'
    HTMLExporter().export(str(ir_path),str(out))
    assert (out.parent/'middle.png').is_file()
    assert (out.parent/'operation.png').is_file()


def test_missing_input_image_does_not_borrow_supporting_image(tmp_path):
    data=graph_payload();data['shapes'][4]['image_path']=None
    path=tmp_path/'graph.json';path.write_text(json.dumps(data))
    guide=build_guide(str(path))
    assert guide.steps[1].input.image_path is None
    document=WizardDocument.from_any(guide_to_dict(guide),source_dir=str(tmp_path))
    out=tmp_path/'missing.pptx';PPTXExporter().export(document,out,ExportOptions())
    found=False
    for slide in Presentation(str(out)).slides:
        text=' '.join(s.text for s in slide.shapes if s.has_text_frame)
        if 'Starting assembly: Middle' in text:
            found=True
            assert 'No image available' in text
            assert not any(s.shape_type==MSO_SHAPE_TYPE.PICTURE for s in slide.shapes)
    assert found


def test_staging_prefers_document_relative_images(tmp_path):
    from src.main.services.image_staging import stage_images
    document=tmp_path/'document';(document/'images').mkdir(parents=True)
    image=document/'images'/'part.png';Image.new('RGB',(20,10),'blue').save(image)
    snapshot={'shapes':[{'image_path':'images/part.png'}]}
    handler=SimpleNamespace(get_full_path=lambda ref:str(tmp_path/'missing'/ref))
    staged=tmp_path/'staged';stage_images(snapshot,staged,handler,document)
    assert (staged/'images'/'part.png').read_bytes()==image.read_bytes()


def test_full_viewport_fallback_still_avoids_large_obstacle():
    view=SimpleNamespace(_measure_shape=lambda shape:None,
                         render_shape=lambda shape:shape,
                         render_bounds=lambda shape:(-10000,-10000,10000,10000),
                         zoom_factor=1,winfo_width=lambda:600,winfo_height=lambda:400,
                         canvasx=lambda x:x,canvasy=lambda y:y)
    diagram=Diagram();diagram.shapes=[ComponentBox(0,0)]
    shape=ComponentBox(0,0)
    shape.x,shape.y=DiagramCanvas.find_free_position(view,shape,diagram)
    assert shape.get_bounds()[0]>10000
