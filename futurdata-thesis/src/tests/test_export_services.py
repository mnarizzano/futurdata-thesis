"""Regression coverage for shared exporters and their MVC service boundary."""

import json
from pathlib import Path
import zipfile

import pytest
from PIL import Image

from src.main.models import ActionCircle, ArrowShape, ComponentBox, Diagram, DiamondStep
from src.main.services.document_export_service import DocumentExportService
from src.main.services.html_exporter.html_exporter import HTMLExporter
from src.main.utils.image_handler import ImageHandler
from src.main.utils.json_exporter import EnhancedJSONExporter


class UnusedRepository:
    def __getattr__(self, name):
        raise AssertionError(f"Document export accessed persistence: {name}")


@pytest.fixture
def export_diagram(tmp_path, monkeypatch):
    handler = ImageHandler(str(tmp_path / "store"))
    image = Path(handler.images_dir) / "part.png"
    Image.new("RGB", (16, 16), "blue").save(image)
    for module in ("document_export_service", "presentation_service"):
        monkeypatch.setattr(
            f"src.main.services.{module}.get_image_handler", lambda: handler
        )
    diagram = Diagram()
    root = ComponentBox(100, 100)
    root.text = "Coffee machine"
    root.properties.update(
        name="Coffee machine", node_type="Root", image_path="images/part.png"
    )
    step = ActionCircle(100, 300)
    step.text = "Remove cover"
    action = DiamondStep(300, 300)
    action.text = action.name = "Unscrew cover"
    action.image_path = "images/part.png"
    leaf = ComponentBox(100, 500)
    leaf.text = "Cover"
    leaf.properties.update(name="Cover", node_type="Leaf", image_path="images/part.png")
    diagram.shapes.extend((root, step, action, leaf))
    for source, target in ((root, action), (action, step), (action, leaf)):
        diagram.shapes.append(ArrowShape(0, 0, source, target))
    return diagram, image


@pytest.mark.parametrize("format_id", ["html", "md", "txt", "docx", "pptx"])
def test_document_exports_preserve_model_and_document_content(
    tmp_path, export_diagram, format_id
):
    diagram, image = export_diagram
    exporter = EnhancedJSONExporter(UnusedRepository())
    before = exporter.serialize_active_diagram(diagram)
    service = DocumentExportService(exporter)
    output = Path(
        service.export(diagram, str(tmp_path / "output" / "guide"), format_id)
    )
    assert output.is_file() and output.stat().st_size > 0
    assert output.suffix == "." + format_id
    after = exporter.serialize_active_diagram(diagram)
    # Export metadata intentionally receives fresh timestamps on each call.
    for snapshot in (before, after):
        for key in ("created", "modified"):
            snapshot["metadata"].pop(key)
    assert after == before
    if format_id in {"html", "md", "txt"}:
        assert (
            output.parent / "images" / "part.png"
        ).read_bytes() == image.read_bytes()
        assert "Coffee machine" in output.read_text(encoding="utf-8")
    else:
        with zipfile.ZipFile(output) as archive:
            prefix = "word/" if format_id == "docx" else "ppt/"
            assert any(
                name.startswith(prefix + "media/") for name in archive.namelist()
            )
            assert any(
                b"Coffee machine" in archive.read(name)
                for name in archive.namelist()
                if name.endswith(".xml")
            )


def test_shared_html_copies_nested_continuation_images(tmp_path):
    source = tmp_path / "source"
    image = source / "images" / "continuation.png"
    image.parent.mkdir(parents=True)
    image.write_bytes(b"nested image")
    ir = {
        "product": {"name": "Nested"},
        "steps": [{"continues_as": [{"image": {"path": "images/continuation.png"}}]}],
    }
    ir_path = source / "guide.json"
    ir_path.write_text(json.dumps(ir), encoding="utf-8")
    output = tmp_path / "output" / "guide.html"
    HTMLExporter().export(str(ir_path), str(output))
    assert (
        output.parent / "images" / "continuation.png"
    ).read_bytes() == image.read_bytes()


def test_pptx_embeds_product_action_and_output_images(tmp_path, export_diagram):
    from pptx import Presentation
    from pptx.enum.shapes import MSO_SHAPE_TYPE

    diagram, image = export_diagram
    expected = set()
    for shape, name, color, size in (
        (diagram.shapes[0], "product.png", "red", (80, 40)),
        (diagram.shapes[1], "instruction.jpg", "green", (30, 60)),
        (diagram.shapes[3], "output.png", "blue", (50, 50)),
    ):
        path = image.parent / name
        Image.new("RGB", size, color).save(path)
        expected.add(path.read_bytes())
        if isinstance(shape, ComponentBox):
            shape.properties["image_path"] = "images/" + name
        else:
            shape.image_path = "images/" + name
    service = DocumentExportService(EnhancedJSONExporter(UnusedRepository()))
    output = service.export(diagram, str(tmp_path / "pictures.pptx"), "pptx")
    # Reopen after the service's temporary image staging directory is removed.
    presentation = Presentation(output)
    pictures = [
        shape
        for slide in presentation.slides
        for shape in slide.shapes
        if shape.shape_type == MSO_SHAPE_TYPE.PICTURE
    ]
    assert expected <= {picture.image.blob for picture in pictures}
    for picture in pictures:
        width, height = picture.image.size
        assert picture.width / picture.height == pytest.approx(width / height, rel=1e-5)
